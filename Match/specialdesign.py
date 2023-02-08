import logging,ast,telebot,redis
import threading,docx,time,json
from flask import Flask,request
from telebot import types,util
from openpyxl import Workbook
from emoji import emojize
#telebot.logger.setLevel(logging.DEBUG) #for debugging telegram api
API_TOKEN = ""#this value will set localy on your host and will obtain from botfather
#this value will set localy on your host and will obtain from botfather
WEBHOOK_HOST = '207.182.136.94'#'<ip/host where the bot is running>'
WEBHOOK_PORT = 8443  # 443, 80, 88 or 8443 (port need to be 'open')
WEBHOOK_LISTEN = '0.0.0.0'  # In some VPS you may need to put here the IP addr
WEBHOOK_URL_BASE = "https://%s:%s" % (WEBHOOK_HOST, WEBHOOK_PORT)
WEBHOOK_URL_PATH = "/%s/" % (API_TOKEN)
server = Flask(__name__)
lines = docx.Document('question.docx')  # Creating word reader object.
fullText = []
redis = redis.Redis(host='localhost', port=6379, db=0)
bot = telebot.TeleBot(API_TOKEN,threaded=True,num_threads=8)

for para in lines.paragraphs:
    fullText.append(para.text)
count,countans,answerlist,questionsdoc,questionlist=0,6,[],{},[]
while countans<len(fullText):
	answerlist.append(int(fullText[countans].rstrip("\n")))
	countans+=7
while count<len(fullText):
	questionlist.append(fullText[count].rstrip("\n"))
	questionsdoc[fullText[count].rstrip("\n")]=[fullText[count+1].rstrip("\n"),fullText[count+2].rstrip("\n"),fullText[count+3].rstrip("\n"),fullText[count+4].rstrip("\n"),fullText[count+5].rstrip("\n")]
	count+=7
crossIcon = u"\u274C"
#bot=telebot.AsyncTeleBot(API_TOKEN) #if you want your bot respond asynchronus
options={emojize(':dart:',use_aliases=True)+' شروع آزمون ':'Fight on',emojize(':trophy:',use_aliases=True)+' امتیاز من ':'My point'}
markup = types.InlineKeyboardMarkup()
for key, value in options.items():
    markup.add(types.InlineKeyboardButton(text=key,
                                          callback_data="['value', '" + value + "', '" + key + "']"))
markup.add(types.InlineKeyboardButton(text=emojize(':globe_with_meridians:',use_aliases=True)+' وبسایت ',
                                      url= "https://salamat.gov.ir"))
mksighup = types.ForceReply(selective=False)
markup2=types.ReplyKeyboardMarkup(resize_keyboard=True)
markup2.row(types.KeyboardButton(u"\u2160"),types.KeyboardButton(u"\u2161"),types.KeyboardButton(u"\u2162"),types.KeyboardButton(u"\u2163"),types.KeyboardButton(u"\u2164"))
sighuplist={}
@bot.message_handler(commands=['end'])
def surrend(message):
	try:
		redis.hset(message.from_user.username,'state','finish')
		bot.reply_to(message, '  شانس خود را از دست دادی {}'.format(message.from_user.first_name),reply_markup=markup)
	except:
		pass
@bot.message_handler(commands=['begin'])
def beginner(message):
	keys=[]
	for data8 in redis.scan_iter():
		keys.append(data8.decode('utf-8'))
	if message.from_user.username in keys:
		if redis.hget(message.from_user.username,'state')==b'finish' :
			bot.reply_to(message, 'شما شانس خود را قبلا استفاده کرده اید\n/home')
		elif redis.hget(message.from_user.username,'state')==b'on match':
			bot.reply_to(message,'شما هم اکنون در آزمون حضور دارید')
			redis.hset(message.from_user.username,'timer',time.time())
			getquestion(message)
		elif redis.hget(message.from_user.username,'state')==b'alive':
			redis.hset(message.from_user.username,'state','on match')
			redis.hset(message.from_user.username,'timer',time.time())
			getquestion(message)
	else:
		bot.reply_to(message, 'لطفا ابتدا ثبت نام کنید \n/start & /begin')

def check(message,num):
    g=questionsdoc[questionlist[sighuplist[message.from_user.username]['questionnumbers'][-2]]][num]
    if (time.time()-float(redis.hget(message.chat.username,'timer')))==45 or (time.time()-float(redis.hget(message.chat.username,'timer')))>45:
        bot.send_message(message.chat.id,'متاسفانه وقت شما تمام شد')
    elif g==questionsdoc[questionlist[sighuplist[message.from_user.username]['questionnumbers'][-2]]][answerlist[sighuplist[message.from_user.username]['questionnumbers'][-2]]-1]:
        data1=int(redis.hget(message.from_user.username,'point'))
        redis.hset(message.from_user.username,'point',data1+20)
    else:
        pass
@bot.callback_query_handler(func=lambda call: True)
def handle_query(call):
    valueFromCallBack = ast.literal_eval(call.data)[1]
    keyFromCallBack = ast.literal_eval(call.data)[2]
    if valueFromCallBack=='My point':
        bot.answer_callback_query(callback_query_id=call.id,
                            show_alert=True,
                            cache_time=3,
                            text="امتیاز شما:{} ".format(redis.hget(call.message.chat.username,'point').decode('utf-8')))
    elif  valueFromCallBack=='Fight on':
        bot.edit_message_text(chat_id=call.message.chat.id,
                              text="Alright you have only 45 sec for each question , 15 question and only one chance \nthese are your commands\nfor surrendering /end \nfor begining /begin ",
                              message_id=call.message.message_id,
                              reply_markup=None,
                              parse_mode='HTML')

def sighup(message):
    if message.reply_to_message != None:
        if message.text=='flow':
            doc = open('output.xlsx', 'rb')
            bot.send_document(message.chat.id, doc)
        elif  len(message.text) != 11:
            msg = bot.reply_to(message, 'لطقا شماره همراه خود را به این صورت وارد کنید: \n ex:09xxxxxxxxx',reply_markup=mksighup)
        else:
            bot.send_message(message.chat.id, " لطفا شماره ملی خود را وارد کنید")
            sighuplist[message.from_user.username]={'firstname':'','lastname':'','meli':'','phonenumber':message.text,'questionnumbers':[0],'state':'alive','point':0,'questionnum':1,'timer':time.time(),'chat_id':message.chat.id}

    elif  message.content_type=='text':
        try:
            if message.text=='flow':
                try:
                  doc = open('output.xlsx', 'rb')
                  bot.send_document(message.chat.id, doc)
                except:
                  print("it is not possible")
                  pass
            if sighuplist[message.from_user.username]['meli']=='':
                sighuplist[message.from_user.username]['meli']=message.text
                bot.reply_to(message, ":لطفا اسم کوچک خود را وارد کنید ")
            elif sighuplist[message.from_user.username]['firstname'] =='':
                sighuplist[message.from_user.username]['firstname']=message.text
                bot.reply_to(message, " :لطفا نام خانوادگی خود را وارد کنید {}".format(message.text))
            elif sighuplist[message.from_user.username]['lastname']=='':
                sighuplist[message.from_user.username]['lastname']=message.text
                bot.reply_to(message, " ثبت نام شما تکمیل شد {} {}".format(sighuplist[message.from_user.username]['firstname'],sighuplist[message.from_user.username]['lastname']))
                bot.send_message(message.chat.id, "Please choose from  this options", reply_markup=markup)
                redis.hset(message.from_user.username,'lastname',sighuplist[message.from_user.username]['lastname'])
                redis.hset(message.from_user.username,'firstname',sighuplist[message.from_user.username]['firstname'])
                redis.hset(message.from_user.username,'phonenumber',sighuplist[message.from_user.username]['phonenumber'])
                redis.hset(message.from_user.username,'meli',sighuplist[message.from_user.username]['meli'])
                redis.hset(message.from_user.username,'state',sighuplist[message.from_user.username]['state'])
                redis.hset(message.from_user.username,'point',sighuplist[message.from_user.username]['point'])
                redis.hset(message.from_user.username,'questionnum',sighuplist[message.from_user.username]['questionnum'])
                redis.hset(message.from_user.username,'chat_id',sighuplist[message.from_user.username]['chat_id'])
                redis.hset(message.from_user.username,'timer',sighuplist[message.from_user.username]['timer'])
            else:
               if redis.hget(message.from_user.username,'state')=='finish':
                  bot.send_message(message.chat.id,'شما شانس خود را قبلا استفاده کرده اید')
                  pass
               elif int(redis.hget(message.from_user.username,'questionnum')) ==15 or int(redis.hget(message.from_user.username,'questionnum'))>15 :
                  bot.send_message(message.chat.id,emojize(':coffee:',use_aliases=True)+'congratulation you have answered to all questions\nand your point is :{} \njust take rest and relax'.format(redis.hget(message.from_user.username,'point').decode('utf-8')),reply_markup=markup)
                  redis.hset(message.from_user.username,'state','finish')
                  del sighuplist[message.from_user.username]
               else:
                 data2=int(redis.hget(message.from_user.username,'questionnum'))+1
                 redis.hset(message.from_user.username,'questionnum',data2)
                 if message.text==u"\u2160":
                   check(message,0)
                   getquestion(message)
                 elif message.text==u"\u2161":
                   check(message,1)
                   getquestion(message)
                 elif message.text==u"\u2162":
                   check(message,2)
                   getquestion(message)
                 elif message.text==u"\u2163":
                   check(message,3)
                   getquestion(message)
                 elif message.text==u"\u2164":
                   check(message,4)
                   getquestion(message)
                 redis.hset(message.from_user.username,'timer',time.time())
        except:
            pass
    else:
        bot.reply_to(message, "wrong input!!!")

def getquestion(message):
    s=questionlist[sighuplist[message.from_user.username]['questionnumbers'][-1]]
    g=questionsdoc[s]
    for text in util.split_string('{}\n{}\n{}\n{}\n{}\n{}'.format("".join([str(redis.hget(message.from_user.username,'questionnum').decode('utf-8')),':',s]),"".join([u"\u2160",":",g[0]]),"".join([u"\u2161",":",g[1]]),"".join([u"\u2162",":",g[2]]),"".join([u"\u2163",":",g[3]]),"".join([u"\u2164",":",g[4]])), 3000):
        bot.send_message(message.chat.id,text=text,reply_markup=markup2)
    number=sighuplist[message.from_user.username]['questionnumbers'][-1]+1
    sighuplist[message.from_user.username]['questionnumbers'].append(number)
    return None

@bot.message_handler(commands=['home', 'start'])
def send_welcome(message):
	try:
		if message.chat.username==None:
			bot.send_message(message.chat.id,"hmmm it seems like that you dont have username account so go to telegram setting and set one to continue!")
		elif redis.hget(message.from_user.username,'firstname')!=None:
			bot.send_message(message.chat.id,  "لطفا از موارد زیر انتخاب کنید", reply_markup=markup)
		else:
			bot.reply_to(message,'hi {} Welcome to home we have examination hear enter your phone number:'.format(message.from_user.first_name),reply_markup=mksighup)
	except:
		pass
@bot.message_handler(func=sighup)
def startmatch(message):
	if redis.hget(message.from_user.username,'firstname') !='':
		bot.reply_to(message, "لطفا از موارد زیر انتخاب کنید", reply_markup=markup)

@bot.message_handler(content_types=['document', 'audio'])
def handle_docs_audio(message):
	pass

@bot.message_handler(commands=['alpha'])
def output(message):
    try:
        workbook = openpyxl.loads_workbook('output.xlsx')
        sheet= workbook.active
    except:
        workbook = Workbook()
        sheet = workbook.active
        pass
    finally:
        print(message.from_user.first_name," is in alpha office")
        bot.reply_to(message,emojize(':open_file_folder:',use_aliases=True)+'Hi mr/mis {} you are in alpha office enter password to receive results'.format(message.from_user.first_name),reply_markup=mksighup)
        for key in redis.scan_iter():
            try:
                value1=redis.hget(key,'firstname')
                value2=redis.hget(key,'lastname')
                value3=redis.hget(key,'phonenumber')
                value4=redis.hget(key,'meli')
                value5=redis.hget(key,'point')
                value6=redis.hget(key,'state')
                new_row = [key, value1,value2,value3,value4,value5,value6]
                sheet.append(new_row)
            except:
                pass
        workbook.save(filename="output.xlsx")


@server.route('/', methods=['GET', 'HEAD'])
def index():
    return ''
@server.route(WEBHOOK_URL_PATH, methods=['POST'])
def webhook():
    if flask.request.headers.get('content-type') == 'application/json':
        json_string = flask.request.get_data().decode('utf-8')
        update = telebot.types.Update.de_json(json_string)
        bot.process_new_updates([update])
        return ''
    else:
        flask.abort(403)
# Set webhook
bot.remove_webhook()
time.sleep(0.2)
bot.set_webhook(url=WEBHOOK_URL_BASE + WEBHOOK_URL_PATH)
if __name__ == "__main__":
    # Start flask server
    server.run(host=WEBHOOK_LISTEN,
            port=WEBHOOK_PORT,
            debug=True)
