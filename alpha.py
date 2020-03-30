import logging,ast,telebot,redis
import threading,docx,time,os
from flask import Flask, request
from telebot import types,util
from random import seed,randint
from openpyxl import Workbook
#telebot.logger.setLevel(logging.DEBUG) #for debugging telegram api
API_TOKEN = os.getenv("API_TOKEN")#this value will set localy on your host and will obtain from botfather
WEBHOOK_LISTEN = '0.0.0.0' #default
server = Flask(__name__)
lines = docx.Document('question.docx')  # Creating word reader object.
fullText = []
for para in lines.paragraphs:
    fullText.append(para.text)
count,countans,answerlist,questionsdoc,questionlist=0,6,[],{},[]
while countans<len(fullText):
	answerlist.append(int(fullText[countans].rstrip("\n")))
	countans+=7
while count<len(fullText):
	questionlist.append(fullText[count].rstrip("\n"))
	questionsdoc[fullText[count].rstrip("\n")]=[fullText[count+1].rstrip("\n"),fullText[count+2].rstrip("\n"),fullText[count+3].rstrip("\n"),fullText[count+4].rstrip("\n"),fullText[count+5].rstrip("\n")]
	count+=7
seed(1)
crossIcon = u"\u274C"
bot = telebot.TeleBot(API_TOKEN,threaded=True,num_threads=6)
#bot=telebot.AsyncTeleBot(API_TOKEN) #if you want your bot respond asynchronus
options={'My point':'My point','Fight on':'Fight on','Home':'Home'}
markup = types.InlineKeyboardMarkup()
for key, value in options.items():
    markup.add(types.InlineKeyboardButton(text=key,
                                          callback_data="['value', '" + value + "', '" + key + "']"))
mksighup = types.ForceReply(selective=False)
markup2=types.ReplyKeyboardMarkup()
markup2.add(types.KeyboardButton(u"\u2160"))
markup2.add(types.KeyboardButton(u"\u2161"))
markup2.add(types.KeyboardButton(u"\u2162"))
markup2.add(types.KeyboardButton(u"\u2163"))
markup2.add(types.KeyboardButton(u"\u2164"))
sighuplist={}
@bot.message_handler(commands=['end'])
def surrend(message):
	try:
		sighuplist[message.from_user.username]['state']='finish'
		bot.reply_to(message, 'oh my friend {} it was so hard for you,i know'.format(message.from_user.first_name),reply_markup=markup)
	except:
		pass
@bot.message_handler(commands=['begin'])
def beginner(message):
	if message.from_user.username in sighuplist.keys():
		if sighuplist[message.from_user.username]['state']=='finish' :
			bot.reply_to(message, 'sorry you lose your chance !!!')
		elif sighuplist[message.from_user.username]['state']=='on match':
			bot.reply_to(message,'you are in match right now!')
			sighuplist[message.from_user.username]['timer']=time.time()
			getquestion(message)
		elif sighuplist[message.from_user.username]['state']=='alive':
			sighuplist[message.from_user.username]['state']='on match'
			sighuplist[message.from_user.username]['timer']=time.time()
			getquestion(message)
	else:
		bot.reply_to(message, 'signup please \n /start & /begin')

def check(message,num):
    g=questionsdoc[questionlist[sighuplist[message.from_user.username]['questionnumbers'][-2]]][num]
    if (time.time()-sighuplist[message.chat.username]['timer'])==45 or (time.time()-sighuplist[message.chat.username]['timer'])>45:
        bot.send_message(message.chat.id,'your time is finished!!!')
    elif g==questionsdoc[questionlist[sighuplist[message.chat.username]['questionnumbers'][-2]]][answerlist[sighuplist[message.chat.username]['questionnumbers'][-2]]-1]:
        sighuplist[message.chat.username]['point']+=20

@bot.callback_query_handler(func=lambda call: True)
def handle_query(call):
    valueFromCallBack = ast.literal_eval(call.data)[1]
    keyFromCallBack = ast.literal_eval(call.data)[2]
    if valueFromCallBack=='My point':
        bot.answer_callback_query(callback_query_id=call.id,
                            show_alert=True,
                            text="Your point is: " + sighuplist[call.message.from_user.username][point])
    elif  valueFromCallBack=='Fight on':
        bot.send_message(call.message.chat.id, "Alright you have only 45 sec for each question , 15 question and only one chance \n this are your commands\nfor surrendering /end \nfor begining /begin ")
    elif valueFromCallBack=='Home':
        send_welcome(call.message)
def sighup(message):
    if message.reply_to_message != None:
        try:
            if message.text=='flow':
                doc = open('output.xlsx', 'rb')
                bot.send_document(message.chat.id, doc)
            elif  len(message.text) != 11:
                msg = bot.reply_to(message, 'wrong input please input your phone number correctly: \n ex:09xxxxxxxxx',reply_markup=mksighup)
            else:
                bot.send_message(message.chat.id, " Ok now enter your first name :")
                sighuplist[message.from_user.username]={'firstname':'','lastname':'','phonenumber':message.text,'questionnumbers':[0],'state':'alive','point':0,'questionnum':0,'timer':time.time(),'chat_id':message.chat.id}
        except:
            pass
    elif  message.content_type=='text':
            if message.text=='flow':
                try:
                  doc = open('output.xlsx', 'rb')
                  bot.send_document(message.chat.id, doc)
                except:
                  print("it is not possible")
                  pass
            if sighuplist[message.from_user.username]['firstname'] =='':
                sighuplist[message.from_user.username]['firstname']=message.text
                bot.reply_to(message, "Good {} enter your lastname:".format(message.text))
            elif sighuplist[message.from_user.username]['lastname']=='':
                sighuplist[message.from_user.username]['lastname']=message.text
                bot.reply_to(message, "Well done {} {} your signup is complete".format(sighuplist[message.from_user.username]['firstname'],sighuplist[message.from_user.username]['lastname']))
                bot.send_message(message.chat.id, "Please choose from  this options", reply_markup=markup)
            elif message.from_user.username in sighuplist.keys():
               if  sighuplist[message.from_user.username]['state']=='finish':
                  bot.send_message(message.chat.id,'sorry you lost your chance')
                  pass
               elif sighuplist[message.from_user.username]['questionnum'] ==15 or sighuplist[message.from_user.username]['questionnum']>15 :
                  bot.send_message(message.chat.id,u"\u26Fe"+'congratulation you have answered to all questions\nand your point is :{} \njust take rest and relax'.format(sighuplist[message.from_user.username]['point']),reply_markup=markup)
                  sighuplist[message.from_user.username]['state']='finish'
               else:
                 sighuplist[message.from_user.username]['questionnum']+=1
                 if message.text==u"\u2160":
                   check(message,0)
                   getquestion(message)
                 elif message.text==u"\u2161":
                   check(message,1)
                   getquestion(message)
                 elif message.text==u"\u2162":
                   check(message,2)
                   getquestion(message)
                 elif message.text==u"\u2163":
                   check(message,3)
                   getquestion(message)
                 elif message.text==u"\u2164":
                   check(message,4)
                   getquestion(message)
                 sighuplist[message.from_user.username]['timer']=time.time()
    else:
        bot.reply_to(message, "wrong input!!!")

def getquestion(message):
    s=questionlist[sighuplist[message.from_user.username]['questionnumbers'][-1]]
    g=questionsdoc[s]
    for text in util.split_string('{}\n{}\n{}\n{}\n{}'.format("".join([str(sighuplist[message.from_user.username]['questionnum']),':',s]),"".join([u"\u2160",":",g[0]]),"".join([u"\u2161",":",g[1]]),"".join([u"\u2162",":",g[2]]),"".join([u"\u2163",":",g[3]])), 3000):
        bot.send_message(message.chat.id,text=text,reply_markup=markup2)
    number=randint(0,len(questionlist)-1)
    while number in sighuplist[message.chat.username]['questionnumbers']:
        number=randint(0,len(questionlist)-1)
        pass
    sighuplist[message.from_user.username]['questionnumbers'].append(number)
    return None

@bot.message_handler(commands=['help', 'start'])
def send_welcome(message):
	try:
		if message.from_user.username==None:
			bot.send_message(message.chat.id,"hmmm it seems like that you dont have username account so go to telegram setting and set one to continue!")
		elif sighuplist.get(message.from_user.username)!=None:
			bot.send_message(message.chat.id, "Please choose from  this options", reply_markup=markup)
		else:
			bot.reply_to(message,'hi {} Welcome to home we have examination hear'.format(message.from_user.first_name))
			bot.reply_to(message, "Enter your phone number:", reply_markup=mksighup)
	except:
		pass
@bot.message_handler(func=sighup)
def startmatch(message):
	if sighuplist[message.from_user.username]['firstname'] !='':
		bot.reply_to(message, "Please choose from options", reply_markup=markup)

@bot.message_handler(content_types=['document', 'audio'])
def handle_docs_audio(message):
	pass

@bot.message_handler(commands=['alpha'])
def output(message):
	try:
		workbook = openpyxl.load_workbook('output.xlsx')
		sheet= workbook.active
	except:
		workbook = Workbook()
		sheet = workbook.active
		pass
	finally:
		print(message.from_user.first_name," is in alpha office")
		bot.reply_to(message, "Hi mr/mis {} you are in alpha office enter password to receive results".format(message.from_user.first_name),reply_markup=mksighup)
		for key , value in sighuplist.items():
			new_row = [key, value['firstname'],value['lastname'],value['phonenumber'],value['point'],value['state']]
			sheet.append(new_row)
		workbook.save(filename="output.xlsx")

@server.route('/' + API_TOKEN, methods=['POST'])
def getMessage():
    bot.process_new_updates([telebot.types.Update.de_json(request.stream.read().decode("utf-8"))])
    return "!", 200
@server.route("/")
def webhook():
    bot.remove_webhook()
    bot.set_webhook(url="https://civilmatch.herokuapp.com/{}".format(API_TOKEN)) #this will create by heroku create command
    return "!", 200
if __name__ == "__main__":
    server.run(host=WEBHOOK_LISTEN, port=int(os.environ.get("PORT", "8443")))
